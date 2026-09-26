#!/usr/bin/env python3
"""LOG_ON strategy deck — 22 widescreen slides, void + amber."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

VOID = RGBColor(0x09, 0x09, 0x0A)
SURFACE = RGBColor(0x11, 0x11, 0x13)
PANEL = RGBColor(0x18, 0x18, 0x20)
AMBER = RGBColor(0xD4, 0x95, 0x2A)
PAPER = RGBColor(0xF4, 0xF4, 0xF8)
HEAD = RGBColor(0xE8, 0xE8, 0xF0)
BODY = RGBColor(0x9A, 0x9A, 0xA8)
MUTED = RGBColor(0x5A, 0x5A, 0x65)
TEAL = RGBColor(0x1E, 0x8C, 0x6A)

W, H = Inches(13.333), Inches(7.5)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color)
    return s


def textbox(slide, l, t, w, h, text, size=18, color=HEAD, bold=False, italic=False, align=PP_ALIGN.LEFT, font="Georgia"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return tb


def add_runs(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT):
    """runs: list of (text, size, color, bold, italic, font)"""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in runs:
        text, size, color, bold, italic, font = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    return tb


def bullets(slide, l, t, w, h, items, size=16, color=BODY):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = "▸  " + item
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return tb


def footer(slide, page, total=22):
    rect(slide, Inches(0), Inches(7.28), W, Inches(0.22), AMBER)
    textbox(slide, Inches(0.5), Inches(7.28), Inches(8), Inches(0.22),
            "LOG_ON  ·  Lagos  ·  Direction infrastructure for African AI",
            size=9, color=VOID, bold=True, font="Calibri")
    textbox(slide, Inches(11.4), Inches(7.28), Inches(1.5), Inches(0.22),
            f"{page:02d} / {total:02d}", size=9, color=VOID, bold=True, align=PP_ALIGN.RIGHT, font="Calibri")


def chrome(slide, kicker):
    rect(slide, Inches(0), Inches(0), W, H, VOID)
    rect(slide, Inches(0), Inches(0), W, Inches(0.08), AMBER)
    textbox(slide, Inches(0.55), Inches(0.22), Inches(6), Inches(0.3),
            "LOG_ON", size=11, color=AMBER, bold=True, font="Calibri")
    textbox(slide, Inches(6.5), Inches(0.22), Inches(6.3), Inches(0.3),
            kicker, size=11, color=MUTED, align=PP_ALIGN.RIGHT, font="Calibri")


def card(slide, l, t, w, h, title, body, accent=AMBER):
    r = rect(slide, l, t, w, h, SURFACE)
    rect(slide, l, t, Inches(0.08), h, accent)
    textbox(slide, l + Inches(0.28), t + Inches(0.18), w - Inches(0.4), Inches(0.4),
            title, size=14, color=AMBER, bold=True, font="Calibri")
    textbox(slide, l + Inches(0.28), t + Inches(0.58), w - Inches(0.45), h - Inches(0.75),
            body, size=13, color=BODY, font="Calibri")


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    n = 0

    def slide():
        nonlocal n
        n += 1
        s = prs.slides.add_slide(blank)
        return s, n

    # 1 Title
    s, p = slide()
    rect(s, 0, 0, W, H, VOID)
    rect(s, 0, 0, Inches(0.18), H, AMBER)
    textbox(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.4),
            "LOG_ON AI SOLUTION AGENCY  ·  LAGOS, NIGERIA  ·  2026",
            size=12, color=AMBER, bold=True, font="Calibri")
    textbox(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.2),
            "Direction infrastructure\nfor African AI",
            size=44, color=PAPER, bold=True, font="Georgia")
    textbox(s, Inches(0.9), Inches(4.6), Inches(10), Inches(1.2),
            "Your AI is executing. Toward what?\nSystems layer · Cognitive tooling · Epistemological frameworks",
            size=18, color=BODY, italic=True, font="Georgia")
    footer(s, p)

    # 2 The question
    s, p = slide()
    chrome(s, "THE OPERATIONAL QUESTION")
    textbox(s, Inches(0.7), Inches(1.3), Inches(12), Inches(2.2),
            "Intelligence without sovereignty\nis sophisticated servitude.",
            size=36, color=PAPER, italic=True, font="Georgia")
    textbox(s, Inches(0.7), Inches(3.7), Inches(11.5), Inches(2.4),
            "The question is not whether the system is smart.\nThe question is: smart in whose image, toward whose ends, at whose cost?\n\nLOG_ON exists because African organisations are consuming someone else’s cognitive infrastructure — and calling it progress.",
            size=18, color=BODY, font="Calibri")
    footer(s, p)

    # 3 Direction Problem
    s, p = slide()
    chrome(s, "SYNTHESIS  ·  WILLIAMS · DEROSIAUX · KOE")
    textbox(s, Inches(0.7), Inches(1.1), Inches(12), Inches(1),
            "The Direction Problem", size=32, color=PAPER, bold=True, font="Georgia")
    textbox(s, Inches(0.7), Inches(2.1), Inches(12), Inches(1),
            "Execution running ahead of understanding. AI makes the error faster and more expensive.",
            size=16, color=BODY, italic=True, font="Georgia")
    card(s, Inches(0.7), Inches(3.3), Inches(3.8), Inches(3.2),
         "WILLIAMS", "Systems thinking must precede analytical execution — or you get flawless work on the wrong outcome. LFA: Impact → Outcomes → Outputs → Inputs.")
    card(s, Inches(4.75), Inches(3.3), Inches(3.8), Inches(3.2),
         "DEROSIAUX", "The agent’s epistemic environment must be designed before tasks are assigned. Highest-impact change is the constraints file, not PATH.")
    card(s, Inches(8.8), Inches(3.3), Inches(3.8), Inches(3.2),
         "KOE", "AI is leverage for skilled humans. It amplifies direction. Trust is the moat. The person directing it is where the magic lies.")
    footer(s, p)

    # 4 Six threads
    s, p = slide()
    chrome(s, "SIX SYNTHESIS THREADS")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.6),
            "One failure, three scales", size=28, color=PAPER, bold=True, font="Georgia")
    items = [
        ("01  Amplification law", "AI amplifies whatever direction it is given — correct or wrong, equally."),
        ("02  Trust as architecture", "Once execution is commoditised, trust is the only durable moat."),
        ("03  Skill ceiling rises", "Lower barrier, higher cap. Advantage concentrates among directors."),
        ("04  Iteration is the system", "Close the loop at CLAUDE.md / LFA, not at the prompt."),
        ("05  Cognitive vs execution", "PAL and SAE change what you can think. ripgrep only changes speed."),
        ("06  Context as value", "Richer context raises the quality ceiling. This is Cultural Bridge, operationalised."),
    ]
    for i, (t, b) in enumerate(items):
        col, row = i % 3, i // 3
        card(s, Inches(0.7 + col * 4.1), Inches(1.9 + row * 2.4), Inches(3.9), Inches(2.2), t, b)
    footer(s, p)

    # 5 LFA
    s, p = slide()
    chrome(s, "REC 1  ·  LOGICAL FRAMEWORK")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.55),
            "LOG_ON’s logic chain", size=28, color=PAPER, bold=True, font="Georgia")
    rows = [
        ("IMPACT", "African cognitive sovereignty. Frontier models reason in African frameworks as first-class capabilities."),
        ("OUTCOMES", "Research funded and cited. Paying clients. Public body of work that is the reference voice."),
        ("OUTPUTS", "PAL, SAE maps, discourse benchmarks, audits, AgentBridge, essays, grant proposals."),
        ("INPUTS", "Founder time. Lagos deployment. 25 years of systems operations. AgentBase in production."),
    ]
    for i, (k, v) in enumerate(rows):
        y = Inches(1.75 + i * 1.2)
        rect(s, Inches(0.7), y, Inches(11.9), Inches(1.08), SURFACE)
        rect(s, Inches(0.7), y, Inches(0.1), Inches(1.08), AMBER)
        textbox(s, Inches(1.0), y + Inches(0.18), Inches(2.2), Inches(0.7), k, size=14, color=AMBER, bold=True, font="Calibri")
        textbox(s, Inches(3.3), y + Inches(0.22), Inches(9), Inches(0.7), v, size=15, color=HEAD, font="Calibri")
    footer(s, p)

    # 6 Mission
    s, p = slide()
    chrome(s, "MISSION")
    textbox(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(3.2),
            "LOG_ON builds the direction infrastructure for African AI — the systems layer, the cognitive tooling, and the epistemological frameworks that make AI work for African contexts, African organisations, and African futures.",
            size=24, color=PAPER, italic=True, font="Georgia")
    footer(s, p)

    # 7 Offers
    s, p = slide()
    chrome(s, "REC 3  ·  OFFER STACK")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.55),
            "Brand is built. The ladder is now purchasable.", size=26, color=PAPER, bold=True, font="Georgia")
    offers = [
        ("TIER 1", "Cultural Bias\nRapid Audit", "₦300K–₦750K\n$200–$500", "48 hours. 5-page report.\nEdTech, fintech, banks, NGOs, government."),
        ("TIER 2", "Agent Environment\nDesign", "₦1.8M–₦4.5M\n$1,200–$3,000", "2–4 weeks. CLAUDE.md architecture.\nCognitive vs execution tooling."),
        ("TIER 3", "African AI Direction\nConsulting", "₦9M–₦21M\n$6,000–$14,000", "Retainer. LFA for the programme.\nObjective-function alignment."),
    ]
    for i, (tier, name, price, body) in enumerate(offers):
        x = Inches(0.7 + i * 4.15)
        rect(s, x, Inches(1.9), Inches(3.95), Inches(4.6), SURFACE)
        rect(s, x, Inches(1.9), Inches(3.95), Inches(0.1), AMBER)
        textbox(s, x + Inches(0.3), Inches(2.15), Inches(3.4), Inches(0.35), tier, size=12, color=AMBER, bold=True, font="Calibri")
        textbox(s, x + Inches(0.3), Inches(2.55), Inches(3.4), Inches(1.2), name, size=20, color=PAPER, bold=True, font="Georgia")
        textbox(s, x + Inches(0.3), Inches(3.85), Inches(3.4), Inches(0.8), price, size=14, color=AMBER, font="Calibri")
        textbox(s, x + Inches(0.3), Inches(4.75), Inches(3.4), Inches(1.4), body, size=13, color=BODY, font="Calibri")
    footer(s, p)

    # 8 Stats
    s, p = slide()
    chrome(s, "THE STAKES")
    stats = [("<1%", "Africa’s share of global\ndata centre capacity"),
             ("83%", "AI startup funding in Q1 2025\nin four African countries"),
             ("2,000+", "African languages — most\ninvisible in frontier LLMs"),
             ("$10B", "AfDB / UNDP AI initiative\ntarget by 2035")]
    for i, (n_, l) in enumerate(stats):
        x = Inches(0.7 + i * 3.15)
        rect(s, x, Inches(2.1), Inches(2.95), Inches(3.6), SURFACE)
        textbox(s, x + Inches(0.2), Inches(2.4), Inches(2.55), Inches(1.3), n_, size=36, color=AMBER, bold=True, font="Georgia")
        textbox(s, x + Inches(0.2), Inches(3.9), Inches(2.55), Inches(1.4), l, size=14, color=BODY, font="Calibri")
    footer(s, p)

    # 9 Whose Intelligence
    s, p = slide()
    chrome(s, "ESSAY NO. 08  ·  ~4,200 WORDS")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.7),
            "Whose Intelligence Is This?", size=28, color=PAPER, bold=True, font="Georgia")
    qs = [
        ("01 Architecture", "Knowledge is not stored. It is generated. Hallucination is a property."),
        ("02 Authorship", "Corpus, RLHF, constitution, operator layer. None of the authors are Africa."),
        ("03 Fragility", "Brittle at the edges. The edge is 1.4 billion people."),
        ("04 Objective", "The reward model is a statement of values, written in mathematics."),
        ("05 African AI", "Not localisation. Computational primitives of African reasoning."),
        ("06 Interface", "Sophisticated misrepresentation is more dangerous than visible incompetence."),
        ("07 Error", "Sycophancy is the design error that most directly harms African users."),
        ("08 Manifesto", "Document. Audit. Benchmark. Compute. Epistemic sovereignty."),
    ]
    for i, (t, b) in enumerate(qs):
        col, row = i % 4, i // 4
        card(s, Inches(0.55 + col * 3.15), Inches(1.95 + row * 2.4), Inches(3.0), Inches(2.2), t, b)
    footer(s, p)

    # 10 Five actions
    s, p = slide()
    chrome(s, "FIVE ACTIONS THAT CANNOT WAIT")
    acts = [
        ("I  Document before you model", "Formalise proverbs, oral genres, deliberation protocols as machine-readable knowledge. The proverb is a causal model."),
        ("II  Audit before you deploy", "SAE feature analysis. Every representation collapse is a documented bias site."),
        ("III  Benchmark before you evaluate", "Proverb completion, communal negotiation, oral narrative coherence — not translated English."),
        ("IV  Compute sovereignty", "You cannot build sovereign AI on infrastructure you rent from San Francisco."),
        ("V  Epistemic sovereignty", "Hardware enables this. It does not constitute it. Think in African frameworks toward African-defined futures."),
    ]
    for i, (t, b) in enumerate(acts):
        y = Inches(1.15 + i * 1.1)
        rect(s, Inches(0.7), y, Inches(11.9), Inches(1.0), SURFACE)
        rect(s, Inches(0.7), y, Inches(0.1), Inches(1.0), AMBER)
        textbox(s, Inches(1.05), y + Inches(0.12), Inches(11.2), Inches(0.32), t, size=16, color=AMBER, bold=True, font="Calibri")
        textbox(s, Inches(1.05), y + Inches(0.46), Inches(11.2), Inches(0.45), b, size=14, color=BODY, font="Calibri")
    footer(s, p)

    # 11 Research programme
    s, p = slide()
    chrome(s, "AFRICAN COGNITIVE AI  ·  BUILD SEQUENCE")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.55),
            "The ontology unlocks the study. The study validates the ontology.", size=24, color=PAPER, bold=True, font="Georgia")
    steps = [
        ("NOW", "PAL v1.0", "Owomoyela + Ehineni + community panels. Unlocks ideas 1, 2, 3, 6, 8."),
        ("FAST", "Matched-pair pivot", "Cultural math-grounding method on discourse. 40 speakers. EMNLP findings."),
        ("CORE", "Study 1 RCT", "n=537 target. Comprehension, 48h retention, effort, trust, modification. CHI / Cognition."),
        ("CORE", "Bergen test", "African proverb vs Western hypothetical. First cross-cultural simulation test."),
        ("LONG", "CHI user study", "Task-based, discourse-native vs standard AI. HCI community entry."),
        ("GRANT", "Schmidt 2026", "Whitebox probes at the translation layer. $947k. 24 months."),
    ]
    for i, (when, title, body) in enumerate(steps):
        col, row = i % 3, i // 3
        card(s, Inches(0.7 + col * 4.1), Inches(1.85 + row * 2.45), Inches(3.9), Inches(2.25),
             f"{when}  ·  {title}", body)
    footer(s, p)

    # 12 Benchmark paradox
    s, p = slide()
    chrome(s, "THE MEASUREMENT FAILURE")
    textbox(s, Inches(0.7), Inches(1.1), Inches(12), Inches(0.7),
            "The Benchmark Paradox", size=32, color=PAPER, bold=True, font="Georgia")
    textbox(s, Inches(0.7), Inches(1.9), Inches(12), Inches(1.1),
            "When a model produces discourse-authentic Yoruba — story-first, community-referenced, proverb-closed — every automated metric penalises it. Human Yoruba speakers prefer it. Correlation is negative.",
            size=16, color=BODY, font="Calibri")
    card(s, Inches(0.7), Inches(3.3), Inches(5.9), Inches(3.2),
         "VERSION A  ·  STANDARD AI", "Definition-first, evidence-cited, recommendation-closed. BLEU and BERTScore reward proximity to English-origin reference. AfroBench reads register shift as error.")
    card(s, Inches(6.85), Inches(3.3), Inches(5.8), Inches(3.2),
         "VERSION B  ·  DISCOURSE-ALIGNED", "Story of Mama Kehinde at Bodija. Maize-roots proverb as closing crystallisation. Human cultural relevance 4.2–4.8 vs 2.0–2.8. The benchmark cannot see this.")
    footer(s, p)

    # 13 Study 1
    s, p = slide()
    chrome(s, "STUDY 1 OF 3")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.7),
            "Does discourse structure determine comprehension?", size=26, color=PAPER, bold=True, font="Georgia")
    bullets(s, Inches(0.7), Inches(1.9), Inches(6), Inches(4.6), [
        "3-group between-subjects. Conditions A / B / C (proverb, story-only, standard AI).",
        "Five outcomes: immediate comprehension, 48-hour retention, cognitive effort, trust, modification.",
        "Powered on the bottleneck: effort d=0.35 → n=155/group, 537 with attrition.",
        "OTE-I: 17-item Oral Tradition Exposure Index.",
        "Pre-committed patterns A / B / C. A null result is useful — it stops five years of unvalidated building.",
        "Zero of 884 African NLP papers (2019–2026) meet all five inclusion criteria.",
    ])
    card(s, Inches(7.2), Inches(1.9), Inches(5.4), Inches(4.6),
         "WHY THIS HAS NOT BEEN DONE", "Disciplinary separation. Cognitive scientist + Yoruba linguist + LLM operator + Nigerian university access have not been in the same room. This study is coordination, not a new capability.")
    footer(s, p)

    # 14 PAL
    s, p = slide()
    chrome(s, "STUDY 2  ·  PROVERB ACTIVATION LIBRARY")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.55),
            "AI can retrieve a proverb. It cannot deploy one.", size=26, color=PAPER, italic=True, font="Georgia")
    dims = [
        "D1 Simulation domain", "D2 Conceptual domain", "D3 Discourse position",
        "D4 Speech act", "D5 Relationship register", "D6 Community specificity", "D7 Failure contexts",
    ]
    for i, d in enumerate(dims):
        y = Inches(1.8 + i * 0.65)
        rect(s, Inches(0.7), y, Inches(6.3), Inches(0.58), SURFACE)
        textbox(s, Inches(0.95), y + Inches(0.12), Inches(5.9), Inches(0.4), d, size=16, color=HEAD, bold=True, font="Calibri")
    card(s, Inches(7.3), Inches(1.8), Inches(5.3), Inches(4.7),
         "CG-CoT ASKED FOR THIS", "June 2025 future work: structured cultural ontologies, retrieval at discourse positions, native speaker panels. PAL is that list, built. JSON-LD + FAISS. Community disagreement is encoded as variance, not error.")
    footer(s, p)

    # 15 SAE
    s, p = slide()
    chrome(s, "CULTURAL BRIDGE TECH  ·  SAE LIMITATIONS")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.7),
            "An SAE that scores zero does not mean no bias.\nIt means the conduit failed.", size=24, color=PAPER, italic=True, font="Georgia")
    fixes = [
        ("Fix 1", "JumpReLU + TopK ensemble. Kill dead latents and shrinkage."),
        ("Fix 2", "Fine-tune SAEs on Yoruba / Hausa / Igbo / Pidgin. Native training beats transfer."),
        ("Fix 3", "Hybrid scoring: SAE + probes + counterfactual deltas."),
        ("Fix 4", "Error probe on reconstruction residual — recover the dark matter."),
        ("Fix 5", "Adversarial variants: translation, code-switch, cultural rephrase."),
        ("Fix 6", "Multi-layer circuit tracing with Gemma Scope 2."),
    ]
    for i, (t, b) in enumerate(fixes):
        col, row = i % 3, i // 3
        card(s, Inches(0.7 + col * 4.1), Inches(2.15 + row * 2.25), Inches(3.9), Inches(2.1), t, b)
    footer(s, p)

    # 16 Schmidt
    s, p = slide()
    chrome(s, "SCHMIDT SCIENCES 2026  ·  $947,343  ·  24 MONTHS")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.7),
            "Detecting deception where it hides", size=28, color=PAPER, bold=True, font="Georgia")
    textbox(s, Inches(0.7), Inches(1.8), Inches(12), Inches(1.2),
            "LLMs map African-language queries to English in the bottom layers, reason in English in the middle, and translate back at the top. Current probes read the output. The deception lives in the middle.",
            size=16, color=BODY, font="Calibri")
    card(s, Inches(0.7), Inches(3.2), Inches(3.9), Inches(3.3),
         "D6  CAPABILITY", "Language-specific neurons show weakness. Output is confident. Whitebox reads the contradiction. Blackbox cannot.")
    card(s, Inches(4.75), Inches(3.2), Inches(3.9), Inches(3.3),
         "D3  CONTEXT", "Middle layers generate English tokens for a Yoruba query. The user believes the interaction is in Yoruba. It is not.")
    card(s, Inches(8.8), Inches(3.2), Inches(3.8), Inches(3.3),
         "WP4  LAGOS", "40 SME operators. Steering as deception reduction, not “human-AI collaboration” — the scope trap the compliance audit closed.")
    footer(s, p)

    # 17 AgentBridge
    s, p = slide()
    chrome(s, "PRODUCT  ·  MIT LICENCE")
    textbox(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.55),
            "AgentBase + AgentBridge Africa", size=28, color=PAPER, bold=True, font="Georgia")
    card(s, Inches(0.7), Inches(1.85), Inches(6.0), Inches(4.6),
         "AGENTBASE", "Lightweight multi-agent orchestration. No framework lock-in. Explicit over magic. Production-first retries and checkpoints. Zero GPU. $6/month VPS. Running HouseHunter, GigPilot, Job Application Command Centre, four client pipelines.")
    card(s, Inches(6.95), Inches(1.85), Inches(5.7), Inches(4.6),
         "AGENTBRIDGE AFRICA", "WhatsApp-first. Yoruba / Pidgin / Igbo config. CRDT offline. Paystack / Flutterwave / Opay. Circle USDC nanopayments for template authors. Pilot: 500 Lagos businesses. Phase 4 of the 30-month research architecture — DRL toggles oral-tradition format.")
    footer(s, p)

    # 18 SMME
    s, p = slide()
    chrome(s, "OLUJIMI ET AL. 2025  ·  66 PAPERS")
    textbox(s, Inches(0.7), Inches(1.1), Inches(12), Inches(1.0),
            "The field is ready for what the African AI series built.\nThe SLR validated the gap. This work fills it.",
            size=22, color=PAPER, italic=True, font="Georgia")
    ideas = [
        ("9.1", "Discourse-native SMME agent — explanations that follow Yoruba / Igbo / Hausa arcs."),
        ("8.8", "First empirical validation of agentic AI in African SMMEs over 6+ months."),
        ("8.8", "Ubuntu as multi-agent alignment — relational, not individualist RLHF."),
        ("8.4", "Proverb-logic prompt library for inventory, finance, HR, customer, advisory."),
        ("7.9", "Ifá as prior multi-agent decision architecture mapped onto LangChain."),
        ("—", "Honest claim map: structural claims supported. Unicorn / roadmap claims still tentative."),
    ]
    for i, (sc, b) in enumerate(ideas):
        col, row = i % 3, i // 3
        card(s, Inches(0.7 + col * 4.1), Inches(2.35 + row * 2.2), Inches(3.9), Inches(2.05),
             f"SCORE  {sc}", b)
    footer(s, p)

    # 19 Waterfall
    s, p = slide()
    chrome(s, "REC 4  ·  CONTENT WATERFALL")
    textbox(s, Inches(0.7), Inches(1.1), Inches(12), Inches(0.6),
            "The essays exist. The multiplier is not running.", size=26, color=PAPER, bold=True, font="Georgia")
    flow = [
        ("ESSAY", "Whose Intelligence Is This?\n4,200 words, sourced."),
        ("YOUTUBE", "15-minute camera script.\nArchitecture → authorship → Africa."),
        ("THREADS", "5 LinkedIn threads from sections.\nOne argument each."),
        ("DMS", "EdTech, fintech, NGO, bank,\ngovernment — 48-hour audit."),
        ("PODCAST", "Outline from the same spine.\nOne recording day."),
        ("DECK", "This file. Shareable with\nfundants, clients, collaborators."),
    ]
    for i, (t, b) in enumerate(flow):
        col, row = i % 3, i // 3
        card(s, Inches(0.7 + col * 4.1), Inches(2.0 + row * 2.35), Inches(3.9), Inches(2.2), t, b)
    footer(s, p)

    # 20 30 day
    s, p = slide()
    chrome(s, "REC 6  ·  30-DAY RUN PLAN")
    textbox(s, Inches(0.7), Inches(1.1), Inches(12), Inches(0.5),
            "Direction first. Then speed.", size=28, color=PAPER, bold=True, font="Georgia")
    weeks = [
        ("WEEK 1", "Publish LFA + mission. Lock three-tier offer copy. Freeze PAL v0.1 schema."),
        ("WEEK 2", "Agent environment audits. 15 outreach DMs. Record YouTube video 1."),
        ("WEEK 3", "LinkedIn waterfall live. Schmidt co-PI outreach. Ibadan / Lagos linguist list."),
        ("WEEK 4", "First paid audit. 10 PAL proverbs annotated. Grant calendar lock."),
    ]
    for i, (t, b) in enumerate(weeks):
        y = Inches(1.85 + i * 1.2)
        rect(s, Inches(0.7), y, Inches(11.9), Inches(1.08), SURFACE)
        rect(s, Inches(0.7), y, Inches(2.2), Inches(1.08), AMBER)
        textbox(s, Inches(0.85), y + Inches(0.32), Inches(1.95), Inches(0.5), t, size=14, color=VOID, bold=True, font="Calibri")
        textbox(s, Inches(3.15), y + Inches(0.28), Inches(9.1), Inches(0.55), b, size=16, color=HEAD, font="Calibri")
    footer(s, p)

    # 21 Moat
    s, p = slide()
    chrome(s, "THE MOAT")
    textbox(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.7),
            "Trust × systems competence × knowing what the agent needs before it breaks.",
            size=20, color=PAPER, italic=True, font="Georgia")
    moat = [
        ("01", "25 years of systems-level operations, upstream to downstream oil and gas."),
        ("02", "Original African cognitive AI research — public, citable, grant-shaped."),
        ("03", "AgentBase already running Lagos client pipelines. Not a slideware framework."),
        ("04", "The only public body of work connecting LFA, epistemic tooling, and African epistemology."),
    ]
    for i, (n_, b) in enumerate(moat):
        y = Inches(2.1 + i * 1.15)
        rect(s, Inches(0.7), y, Inches(11.9), Inches(1.05), SURFACE)
        textbox(s, Inches(1.0), y + Inches(0.28), Inches(1.2), Inches(0.5), n_, size=22, color=AMBER, bold=True, font="Georgia")
        textbox(s, Inches(2.4), y + Inches(0.32), Inches(9.8), Inches(0.5), b, size=16, color=HEAD, font="Calibri")
    footer(s, p)

    # 22 Close
    s, p = slide()
    rect(s, 0, 0, W, H, VOID)
    rect(s, 0, 0, Inches(0.18), H, AMBER)
    textbox(s, Inches(0.9), Inches(1.6), Inches(11), Inches(0.4),
            "LOG_ON  ·  LAGOS  ·  2026", size=12, color=AMBER, bold=True, font="Calibri")
    textbox(s, Inches(0.9), Inches(2.15), Inches(11.5), Inches(1.8),
            "Are you a cognitive tenant?", size=36, color=PAPER, italic=True, font="Georgia")
    textbox(s, Inches(0.9), Inches(4.1), Inches(11), Inches(1.4),
            "Cultural Bias Rapid Audit · 48 hours · 5-page report\nlogonthepage@gmail.com",
            size=18, color=BODY, font="Calibri")
    textbox(s, Inches(0.9), Inches(5.8), Inches(11), Inches(0.6),
            "The papers are on the site. The files download. The work is ready to run.",
            size=14, color=MUTED, font="Calibri")
    footer(s, p)

    out_dir = Path("/workspace/public/papers")
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "LOG_ON-Direction-Infrastructure.pptx"
    prs.save(out)
    art = Path("/workspace/artifacts")
    art.mkdir(exist_ok=True)
    prs.save(art / "LOG_ON-Direction-Infrastructure.pptx")
    print("slides", n, "bytes", out.stat().st_size)


if __name__ == "__main__":
    build()
